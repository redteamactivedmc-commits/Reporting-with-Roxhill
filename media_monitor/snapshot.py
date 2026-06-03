"""
Command 1 - Snapshot
Creates a Word document with:
  - Client logo + Active logo header
  - Article metadata table (Publication, Headline, Date, Link)
  - Full-page screenshot of the article URL
  - Saved as: ADMC_<Client>_<Magazine>_<DateOfPublishing>.docx
"""

import os
import tempfile
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from rich.console import Console
from rich.panel import Panel

from media_monitor.config import SNAPSHOTS_DIR
from media_monitor.utils.file_naming import snapshot_filename
from media_monitor.utils.logo_handler import add_logo_header
from media_monitor.utils.browser import capture_url_screenshot

console = Console()


def _add_metadata_table(doc: Document, metadata: dict) -> None:
    """
    Insert the 4-row info table:
    | Publication | <value> |
    | Headline    | <value> |
    | Date        | <value> |
    | Link        | <value> |
    """
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"

    labels = ["Publication", "Headline", "Date", "Link"]
    keys   = ["publication", "headline", "date", "url"]
    colors = {
        "Publication": RGBColor(0x00, 0x00, 0x00),
        "Headline":    RGBColor(0x00, 0x00, 0x00),
        "Date":        RGBColor(0x00, 0x00, 0x00),
        "Link":        RGBColor(0x00, 0x5B, 0xB5),
    }

    for i, (label, key) in enumerate(zip(labels, keys)):
        # Label cell (bold, light blue background)
        label_cell = table.cell(i, 0)
        label_cell.width = Inches(1.3)
        lp = label_cell.paragraphs[0]
        run = lp.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        label_cell._tc.get_or_add_tcPr().append(
            _shading_element("D9EAF7")
        )

        # Value cell
        value_cell = table.cell(i, 1)
        vp = value_cell.paragraphs[0]
        vrun = vp.add_run(str(metadata.get(key, "")))
        vrun.font.size = Pt(10)
        vrun.font.color.rgb = colors[label]
        if key == "url":
            vrun.font.underline = True

    doc.add_paragraph()   # spacing


def _shading_element(hex_color: str):
    """Return an XML shading element for table cell backgrounds."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    return shd


def _add_screenshot(doc: Document, screenshot_path: str) -> None:
    """Insert the screenshot image into the document."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    # Fit to page width (approx 6 inches for A4 with margins)
    run.add_picture(screenshot_path, width=Inches(6.3))


def create_snapshot(
    url: str,
    client: str,
    magazine: str,
    headline: str = "",
    date_str: str = "",
    publication: str = "",
) -> str:
    """
    Main entry point for Command 1.
    Returns the path of the generated .docx file.
    """
    console.print(
        Panel(
            f"[bold cyan]Creating Snapshot[/bold cyan]\n"
            f"Client   : [yellow]{client}[/yellow]\n"
            f"Magazine : [yellow]{magazine}[/yellow]\n"
            f"URL      : {url}",
            title="ADMC Media Monitor",
            border_style="blue",
        )
    )

    # Step 1: Determine date
    date_str = date_str or datetime.today().strftime("%d %B %Y")

    # Step 2: Take screenshot
    console.print("[cyan]  -> Capturing screenshot...[/cyan]")
    tmp_img = os.path.join(tempfile.gettempdir(), "admc_snapshot_tmp.png")
    capture_url_screenshot(url, tmp_img)
    console.print(f"[green]  Done: Screenshot saved to temp: {tmp_img}[/green]")

    # Step 3: Build Word doc
    console.print("[cyan]  -> Building Word document...[/cyan]")
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin   = Inches(0.7)
        section.right_margin  = Inches(0.7)

    # Header logos
    add_logo_header(doc, client)

    # Metadata table
    _add_metadata_table(doc, {
        "publication": publication or magazine,
        "headline":    headline,
        "date":        date_str,
        "url":         url,
    })

    # Screenshot
    _add_screenshot(doc, tmp_img)

    # Step 4: Save
    out_path = snapshot_filename(client, magazine, date_str)
    doc.save(out_path)

    console.print(
        f"\n[bold green]Snapshot created:[/bold green] [underline]{out_path}[/underline]"
    )
    return out_path
