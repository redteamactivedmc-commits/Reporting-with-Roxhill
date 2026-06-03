"""
Logo handler - inserts client logo (left) and Active logo (right)
at the top of every Word document.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from media_monitor.config import ACTIVE_LOGO_PATH, LOGOS_DIR


def _find_client_logo(client_name: str) -> str | None:
    """
    Look for a logo file matching the client name in LOGOS_DIR.
    Accepts .png / .jpg / .jpeg / .gif
    """
    safe_name = client_name.lower().replace(" ", "_")
    for ext in [".png", ".jpg", ".jpeg", ".gif"]:
        candidate = os.path.join(LOGOS_DIR, f"{safe_name}{ext}")
        if os.path.exists(candidate):
            return candidate
    # Try original casing
    for ext in [".png", ".jpg", ".jpeg", ".gif"]:
        candidate = os.path.join(LOGOS_DIR, f"{client_name}{ext}")
        if os.path.exists(candidate):
            return candidate
    return None


def add_logo_header(doc: Document, client_name: str) -> None:
    """
    Adds a two-column header row:
      LEFT  -> Client logo  (or placeholder text)
      RIGHT -> Active/ADMC logo
    """
    # Use a 1-row, 2-column table for alignment
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    # Remove borders for a clean look
    for cell in table.rows[0].cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "none")
            tcBorders.append(border)
        tcPr.append(tcBorders)

    left_cell  = table.cell(0, 0)
    right_cell = table.cell(0, 1)

    # LEFT: Client Logo
    left_para = left_cell.paragraphs[0]
    left_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = left_para.add_run()

    client_logo = _find_client_logo(client_name)
    if client_logo:
        run.add_picture(client_logo, width=Inches(1.8))
    else:
        # Placeholder text styled as bold brand name
        run.text = client_name
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    # RIGHT: Active Logo
    right_para = right_cell.paragraphs[0]
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run2 = right_para.add_run()

    if os.path.exists(ACTIVE_LOGO_PATH):
        run2.add_picture(ACTIVE_LOGO_PATH, width=Inches(1.4))
    else:
        run2.text = "ACTIVE | DIGITAL . MARKETING . COMMUNICATIONS"
        run2.bold = True
        run2.font.size = Pt(8)
        run2.font.color.rgb = RGBColor(0x00, 0xAE, 0xEF)

    # Add a separator line after the header
    sep = doc.add_paragraph()
    sep.paragraph_format.space_after  = Pt(4)
    sep.paragraph_format.space_before = Pt(4)
    pPr  = sep._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "00AEEF")
    pBdr.append(bottom)
    pPr.append(pBdr)
