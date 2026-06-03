"""
Excel helper - create / update the ADMC tracker workbook.
"""

import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

from media_monitor.config import TRACKER_COLUMNS


# Style constants
HEADER_FILL   = PatternFill("solid", fgColor="FFD700")   # Gold
ALT_FILL      = PatternFill("solid", fgColor="F2F2F2")   # Light grey
HEADER_FONT   = Font(bold=True, color="000000", size=10)
CELL_FONT     = Font(size=9)
CENTER        = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT          = Alignment(horizontal="left",   vertical="center", wrap_text=True)
THIN          = Side(style="thin", color="CCCCCC")
THIN_BORDER   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COL_WIDTHS = {
    "S/N":               5,
    "Date (DD/MM/YYYY)": 15,
    "Publication":       22,
    "Headline":          48,
    "Media Type":        12,
    "Language":          10,
    "Spokesperson":      18,
    "Source":            32,
    "Reach":             12,
    "Tier":               8,
    "Print/Online":      12,
    "For Pivot":         10,
}


def _apply_header_row(ws, columns: list[str]) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = CENTER
        cell.border    = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = (
            COL_WIDTHS.get(col_name, 14)
        )
    ws.row_dimensions[1].height = 22


def _style_data_row(ws, row_idx: int, num_cols: int) -> None:
    fill = ALT_FILL if row_idx % 2 == 0 else PatternFill()
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font      = CELL_FONT
        cell.border    = THIN_BORDER
        cell.fill      = fill
        cell.alignment = CENTER if col_idx in [1, 5, 6, 9, 10, 11] else LEFT


def create_or_load_tracker(filepath: str) -> tuple:
    """
    Return (workbook, worksheet).
    Creates a new styled workbook if file doesn't exist.
    """
    if os.path.exists(filepath):
        wb = load_workbook(filepath)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Coverage Tracker"
        ws.freeze_panes = "A2"
        _apply_header_row(ws, TRACKER_COLUMNS)
        wb.save(filepath)
    return wb, ws


def append_tracker_row(ws, data: dict, sn: int) -> None:
    """
    Append a single data row to the tracker worksheet.
    `data` keys should match TRACKER_COLUMNS.
    """
    row_idx = sn + 1   # +1 because row 1 is header
    row_data = [
        sn,
        data.get("Date (DD/MM/YYYY)", ""),
        data.get("Publication", ""),
        data.get("Headline", ""),
        data.get("Media Type", ""),
        data.get("Language", ""),
        data.get("Spokesperson", ""),
        data.get("Source", ""),
        data.get("Reach", ""),
        data.get("Tier", ""),
        data.get("Print/Online", ""),
        data.get("For Pivot", ""),
    ]
    for col_idx, value in enumerate(row_data, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)
    _style_data_row(ws, row_idx, len(TRACKER_COLUMNS))


def next_sn(ws) -> int:
    """Return the next available serial number (max existing + 1)."""
    max_sn = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and isinstance(row[0], int):
            max_sn = max(max_sn, row[0])
    return max_sn + 1
